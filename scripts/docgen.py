"""Shared DOCX/PDF renderers for the manuscript and status reports.

Block types (tuples):
  ("h1", text) ("h2", text) ("p", text) ("bullet", text) ("pagebreak",)
  ("table", caption, header, rows[, col_weights])
  ("fig", path, caption, width_inches)
  ("tag", change_id | None)   -> subsequent blocks belong to that change
Inline markup: **bold**. With highlight=True, tagged blocks are marked
"[Cn]" and shaded yellow so that changes against the base paper stand out.
"""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from PIL import Image as PILImage
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import cm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Image, KeepTogether, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

YELLOW = colors.Color(1, 0.95, 0.55)


def _runs(par, text, size=None, bold_all=False, highlight=False, font="Times New Roman"):
    for tok in re.split(r"(\*\*[^*]+\*\*)", text):
        if not tok:
            continue
        b = tok.startswith("**")
        r = par.add_run(tok.strip("*") if b else tok)
        r.bold = b or bold_all
        r.font.name = font
        r._element.rPr.rFonts.set(qn("w:eastAsia"), font)
        if size:
            r.font.size = Pt(size)
        if highlight:
            r.font.highlight_color = WD_COLOR_INDEX.YELLOW


def _page_number(section):
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    for kind, txt in (("begin", None), (None, "PAGE"), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = txt
        run._r.append(el)


def _borders(table):
    props = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "bottom", "insideH"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "6")
        el.set(qn("w:color"), "000000")
        borders.append(el)
    props.append(borders)


def build_docx(blocks, path: Path, meta: dict, highlight=False, line_numbers=True):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.27), Inches(11.69)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Inches(1))
    if line_numbers:
        ln = OxmlElement("w:lnNumType")
        ln.set(qn("w:countBy"), "1")
        ln.set(qn("w:restart"), "continuous")
        sec._sectPr.append(ln)
    _page_number(sec)
    st = doc.styles["Normal"]
    st.font.name, st.font.size = "Times New Roman", Pt(12)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    st.paragraph_format.line_spacing = 1.5
    st.paragraph_format.space_after = Pt(6)
    doc.core_properties.title = meta["title"]
    doc.core_properties.author = ""

    def para(text="", bold=False, size=12, align=None, hl=False, after=6, spacing=None):
        p = doc.add_paragraph()
        if align == "c":
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif align == "j":
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.space_after = Pt(after)
        if spacing:
            p.paragraph_format.line_spacing = spacing
        if text:
            _runs(p, text, size, bold, hl)
        return p

    para(meta["title"], bold=True, size=16, align="c", after=12)
    for line in meta.get("author_lines", []):
        para(line, align="c", after=2)
    if meta.get("subtitle"):
        para(meta["subtitle"], align="c", after=10)
    if meta.get("abstract"):
        para("Abstract", bold=True, size=13)
        for head, txt in meta["abstract"]:
            para(f"**{head}:** {txt}", size=11, align="j", spacing=1.15)
    if meta.get("keywords"):
        para(f"**Keywords:** {meta['keywords']}", size=11, after=12)
    if meta.get("preface"):
        for blk in meta["preface"]:
            _render_docx(doc, blk, para, False)
    if meta.get("abstract") or meta.get("preface"):
        doc.add_page_break()

    tag = None
    for blk in blocks:
        if blk[0] == "tag":
            tag = blk[1]
            continue
        _render_docx(doc, blk, para, highlight and tag is not None, tag if highlight else None)

    if meta.get("declarations"):
        para("Declarations", bold=True, size=14)
        for head, txt in meta["declarations"]:
            para(f"**{head}.** {txt}", align="j")
    if meta.get("references"):
        para("References", bold=True, size=14)
        for i, ref in enumerate(meta["references"], 1):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.35)
            p.paragraph_format.first_line_indent = Inches(-0.35)
            p.paragraph_format.line_spacing = 1.15
            _runs(p, f"{i}. {ref}", 11)
    doc.save(path)


def _render_docx(doc, blk, para, hl, tag=None):
    k = blk[0]
    mark = f"[{tag}] " if (hl and tag) else ""
    if k == "h1":
        h = para(mark + blk[1], bold=True, size=14, hl=hl)
        h.paragraph_format.keep_with_next = True
        h.paragraph_format.space_before = Pt(12)
    elif k == "h2":
        h = para(mark + blk[1], bold=True, size=12, hl=hl)
        h.paragraph_format.keep_with_next = True
    elif k == "p":
        para(mark + blk[1], align="j", hl=hl)
    elif k == "bullet":
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.line_spacing = 1.5
        _runs(p, mark + blk[1], highlight=hl)
    elif k == "pagebreak":
        doc.add_page_break()
    elif k == "table":
        cap, head, rows = blk[1], blk[2], blk[3]
        c = para(mark + cap, bold=True, size=11, hl=hl, after=3)
        c.paragraph_format.keep_with_next = True
        t = doc.add_table(rows=1, cols=len(head))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        _borders(t)
        for i, h in enumerate(head):
            cell = t.rows[0].cells[i]
            cell.paragraphs[0].paragraph_format.line_spacing = 1.0
            _runs(cell.paragraphs[0], str(h), 9, True)
        for row in rows:
            cells = t.add_row().cells
            for i, v in enumerate(row):
                cells[i].paragraphs[0].paragraph_format.line_spacing = 1.0
                _runs(cells[i].paragraphs[0], str(v), 9)
        para("", after=6)
    elif k == "fig":
        _, path, cap, width = blk
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        iw, ih = PILImage.open(path).size
        p.add_run().add_picture(str(path), width=Inches(min(width, 6.2, 8.0 * iw / ih)))
        p.paragraph_format.keep_with_next = True
        para(mark + cap, size=11, align="j", hl=hl)


def build_pdf(blocks, path: Path, meta: dict, highlight=False):
    fonts = Path("/System/Library/Fonts/Supplemental")
    for name, f in (("TNR", "Times New Roman.ttf"), ("TNR-B", "Times New Roman Bold.ttf"),
                    ("TNR-I", "Times New Roman Italic.ttf"), ("TNR-BI", "Times New Roman Bold Italic.ttf")):
        try:
            pdfmetrics.getFont(name)
        except KeyError:
            pdfmetrics.registerFont(TTFont(name, str(fonts / f)))
    pdfmetrics.registerFontFamily("TNR", normal="TNR", bold="TNR-B", italic="TNR-I", boldItalic="TNR-BI")
    base = ParagraphStyle("b", fontName="TNR", fontSize=11, leading=16.5, alignment=TA_JUSTIFY, spaceAfter=6)
    S = {
        "title": ParagraphStyle("t", parent=base, fontName="TNR-B", fontSize=16, leading=21, alignment=TA_CENTER, spaceAfter=12),
        "c": ParagraphStyle("c", parent=base, alignment=TA_CENTER, spaceAfter=3),
        "h1": ParagraphStyle("h1", parent=base, fontName="TNR-B", fontSize=13.5, leading=18, spaceBefore=12, alignment=0, keepWithNext=1),
        "h2": ParagraphStyle("h2", parent=base, fontName="TNR-B", fontSize=11.5, leading=16, alignment=0, keepWithNext=1),
        "abs": ParagraphStyle("a", parent=base, fontSize=10.5, leading=14),
        "cap": ParagraphStyle("cap", parent=base, fontSize=10, leading=13),
        "capb": ParagraphStyle("capb", parent=base, fontName="TNR-B", fontSize=10, leading=13, alignment=0, keepWithNext=1),
        "cell": ParagraphStyle("cell", parent=base, fontSize=8, leading=10, alignment=0, spaceAfter=0),
        "cellb": ParagraphStyle("cellb", parent=base, fontName="TNR-B", fontSize=8, leading=10, alignment=0, spaceAfter=0),
        "ref": ParagraphStyle("ref", parent=base, fontSize=10, leading=13, leftIndent=18, firstLineIndent=-18),
        "bullet": ParagraphStyle("bu", parent=base, leftIndent=18, bulletIndent=6),
    }
    hl = lambda st: ParagraphStyle(st.name + "h", parent=st, backColor=YELLOW)
    esc = lambda s: s.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
    md = lambda s: re.sub(r"\*\*([^*]+)\*\*", r"<b>\1</b>", esc(s))
    width = A4[0] - 5 * cm

    story = [Paragraph(esc(meta["title"]), S["title"])]
    for line in meta.get("author_lines", []):
        story.append(Paragraph(esc(line), S["c"]))
    if meta.get("subtitle"):
        story.append(Paragraph(esc(meta["subtitle"]), S["c"]))
    story.append(Spacer(1, 10))
    if meta.get("abstract"):
        story.append(Paragraph("<b>Abstract</b>", S["h1"]))
        for head, txt in meta["abstract"]:
            story.append(Paragraph(f"<b>{head}:</b> {esc(txt)}", S["abs"]))
    if meta.get("keywords"):
        story.extend([Spacer(1, 6), Paragraph(f"<b>Keywords:</b> {esc(meta['keywords'])}", S["abs"])])

    def emit(blk, tag, on):
        k = blk[0]
        mark = f"<b>[{tag}]</b> " if (on and tag) else ""
        sty = (lambda s: hl(s)) if on else (lambda s: s)
        if k in ("h1", "h2"):
            story.append(Paragraph(mark + esc(blk[1]), sty(S[k])))
        elif k == "p":
            story.append(Paragraph(mark + md(blk[1]), sty(base)))
        elif k == "bullet":
            story.append(Paragraph(mark + md(blk[1]), sty(S["bullet"]), bulletText="•"))
        elif k == "pagebreak":
            story.append(PageBreak())
        elif k == "table":
            cap, head, rows = blk[1], blk[2], blk[3]
            weights = blk[4] if len(blk) > 4 else None
            data = [[Paragraph(esc(str(h)), S["cellb"]) for h in head]] + \
                   [[Paragraph(md(str(v)), S["cell"]) for v in r] for r in rows]
            n = len(head)
            if weights:
                cw = [width * w / sum(weights) for w in weights]
            elif n == 2:
                cw = [width * 0.38, width * 0.62]
            else:
                cw = [width / n] * n
            tbl = Table(data, colWidths=cw, repeatRows=1)
            tbl.setStyle(TableStyle([("LINEABOVE", (0, 0), (-1, 0), 0.8, colors.black),
                                     ("LINEBELOW", (0, 0), (-1, 0), 0.5, colors.black),
                                     ("LINEBELOW", (0, -1), (-1, -1), 0.8, colors.black),
                                     ("VALIGN", (0, 0), (-1, -1), "TOP"),
                                     ("TOPPADDING", (0, 0), (-1, -1), 2), ("BOTTOMPADDING", (0, 0), (-1, -1), 2)]
                                    + ([("BACKGROUND", (0, 1), (-1, -1), colors.Color(1, 0.98, 0.8))] if on else [])))
            story.extend([Paragraph(mark + esc(cap), sty(S["capb"])), tbl, Spacer(1, 10)])
        elif k == "fig":
            _, p, cap, wid = blk
            w, h = PILImage.open(p).size
            tw = min(wid * 72 * 0.9, width, 600 * w / h)
            story.append(KeepTogether([Image(str(p), width=tw, height=tw * h / w), Spacer(1, 4),
                                       Paragraph(mark + esc(cap), sty(S["cap"])), Spacer(1, 8)]))

    for blk in meta.get("preface", []):
        emit(blk, None, False)
    if meta.get("abstract") or meta.get("preface"):
        story.append(PageBreak())
    tag = None
    for blk in blocks:
        if blk[0] == "tag":
            tag = blk[1]
            continue
        emit(blk, tag, highlight and tag is not None)
    if meta.get("declarations"):
        story.append(Paragraph("Declarations", S["h1"]))
        for head, txt in meta["declarations"]:
            story.append(Paragraph(f"<b>{esc(head)}.</b> {esc(txt)}", base))
    if meta.get("references"):
        story.append(Paragraph("References", S["h1"]))
        for i, ref in enumerate(meta["references"], 1):
            story.append(Paragraph(f"{i}. {esc(ref)}", S["ref"]))

    def footer(canvas, doc_):
        canvas.setFont("TNR", 9)
        canvas.drawCentredString(A4[0] / 2, 1.2 * cm, str(doc_.page))

    SimpleDocTemplate(str(path), pagesize=A4, leftMargin=2.5 * cm, rightMargin=2.5 * cm, topMargin=2.3 * cm,
                      bottomMargin=2.3 * cm, title=meta["title"], author="").build(
        story, onFirstPage=footer, onLaterPages=footer)
